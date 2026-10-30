# -*- coding: utf-8 -*-
"""DB 视图/表导出工具（通用版）

把任意 SQLite 库里的视图（可选表）分别导出成 Excel(.xlsx) 或 CSV。

用法：
    1) 双击 DB视图导出.exe，把 .db（或 .lnk、或含 .db 的文件夹）拖进窗口；
    2) 也可以直接把 .db 拖到 exe 图标上；
    3) 命令行：
       DB视图导出.exe --db D:\\x.db --yes                 # 全导出，xlsx
       DB视图导出.exe --db D:\\x.db --yes --format csv
       DB视图导出.exe --db D:\\x.db --objects tables --yes
       DB视图导出.exe --db D:\\x.db --optimize --yes      # 临时副本上优化后导出
       DB视图导出.exe --db D:\\x.db --optimize-inplace    # 直接改这个库（需输 YES）

安全约定：
    · 默认**只读**打开（mode=ro&immutable=1），不改原库、不产生 -wal/-shm；
    · --optimize 只改**临时副本**，导出结束副本删除，原文件一个字节都不动；
    · --optimize-inplace 会真的改文件（建索引/ANALYZE），必须先输 YES，
      导出结束自动删除自己建的索引；**请只用备份库**。
"""
import argparse
import csv
import os
import shutil
import sqlite3
import sys
import tempfile
import threading
import time
from datetime import datetime

try:
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    HAS_XLSX = True
except Exception:
    HAS_XLSX = False

LONG_CELL = 32000          # Excel 单元格上限 32767，留点余量
PRESET_BANLIST = ["玩家总览", "玩家列表", "BFBAN封禁记录", "BFEAC封禁记录",
                  "TVBOT封禁记录", "TVBOT群组总览", "大致统计", "变更记录"]


# ---------------------------------------------------------------- 通用小工具
def desktop_dir():
    for name in ("Desktop", "桌面"):
        p = os.path.join(os.path.expanduser("~"), name)
        if os.path.isdir(p):
            return p
    p = os.path.join(os.environ.get("USERPROFILE", ""), "Desktop")
    if os.path.isdir(p):
        return p
    return os.getcwd()


def clean_path(text):
    text = (text or "").strip()
    if len(text) >= 2 and text[0] == text[-1] and text[0] in "\"'":
        text = text[1:-1]
    if len(text) >= 2 and text[0] == "&" and text[-1] == "&":
        text = text[1:-1]
    return text.strip()


def resolve_lnk(path):
    try:
        import subprocess
        cmd = ("(New-Object -ComObject WScript.Shell)"
               ".CreateShortcut('%s').TargetPath" % path.replace("'", "''"))
        out = subprocess.run(
            ["powershell", "-NoProfile", "-NonInteractive", "-Command", cmd],
            capture_output=True, text=True, timeout=15,
            encoding="utf-8", errors="replace")
        lines = [ln.strip() for ln in (out.stdout or "").splitlines() if ln.strip()]
        return lines[-1] if lines else None
    except Exception:
        return None


def pick_db_in(folder):
    best, best_t = None, -1
    try:
        for name in os.listdir(folder):
            if not name.lower().endswith(".db"):
                continue
            p = os.path.join(folder, name)
            if os.path.isfile(p) and os.path.getmtime(p) > best_t:
                best, best_t = p, os.path.getmtime(p)
    except Exception:
        return None
    return best


def resolve_input(path):
    path = os.path.abspath(path)
    if os.path.isdir(path):
        db = pick_db_in(path)
        return db, (None if db else "文件夹里没找到 .db")
    if not os.path.isfile(path):
        return None, "路径不存在"
    if path.lower().endswith(".lnk"):
        target = resolve_lnk(path)
        if not target:
            return None, "解析快捷方式失败"
        if os.path.isdir(target):
            db = pick_db_in(target)
            return db, (None if db else "快捷方式指向的文件夹里没有 .db")
        return (target, None) if os.path.isfile(target) else (None, "目标不存在")
    return path, None


def uri_of(path, readonly=True, immutable=True):
    base = "file:%s" % path.replace("\\", "/").replace("#", "%23")
    if readonly:
        return base + ("?mode=ro&immutable=1" if immutable else "?mode=ro")
    return base + "?mode=rw"


def disp_width(text):
    import unicodedata
    return sum(2 if unicodedata.east_asian_width(ch) in ("W", "F") else 1
               for ch in str(text))


def cell_value(v):
    if v is None:
        return None
    if isinstance(v, (bytes, bytearray)):
        s = v.hex()
    else:
        s = v if isinstance(v, str) else str(v)
    if len(s) > LONG_CELL:
        return s[:LONG_CELL] + "…（已截断）"
    return s


def _hms(sec):
    sec = int(max(0, sec))
    if sec >= 3600:
        return "%d:%02d:%02d" % (sec // 3600, (sec % 3600) // 60, sec % 60)
    return "%d:%02d" % (sec // 60, sec % 60)


# ---------------------------------------------------------------- 打开/拷贝/优化
def open_ro(path):
    last = None
    attempts = [(uri_of(path, True, True), True),
                (uri_of(path, True, False), True),
                (path, False)]
    for target, is_uri in attempts:
        try:
            conn = sqlite3.connect(target, uri=is_uri, timeout=15.0)
            conn.row_factory = sqlite3.Row
            conn.execute("PRAGMA query_only=1")
            for pragma in ("PRAGMA temp_store=MEMORY",
                           "PRAGMA cache_size=-131072",
                           "PRAGMA mmap_size=268435456"):
                try:
                    conn.execute(pragma)
                except Exception:
                    pass
            conn.execute("SELECT count(*) FROM sqlite_master").fetchone()
            return conn, None
        except Exception as e:                       # noqa: BLE001
            last = e
    return None, last


def open_rw(path):
    conn = sqlite3.connect(uri_of(path, False), uri=True, timeout=15.0)
    conn.row_factory = sqlite3.Row
    return conn


def side_files(path):
    return [path + "-wal", path + "-shm"]


def snapshot_side(path):
    return {p: os.path.exists(p) for p in side_files(path)}


def cleanup_side(path, before):
    removed = []
    for p in side_files(path):
        if before.get(p):
            continue
        if os.path.exists(p):
            try:
                os.remove(p)
                removed.append(os.path.basename(p))
            except OSError:
                pass
    return removed


class CopyCancelled(Exception):
    pass


def _fmt_mb(n):
    return "%.0f MB" % (n / 1048576.0) if n else "0 MB"


def _kbhit_cancel():
    """复制期间按 n 取消（只在真控制台里生效）。"""
    try:
        import msvcrt
        if msvcrt.kbhit():
            ch = msvcrt.getwch()
            if ch and ch.lower() == "n":
                return True
    except Exception:
        pass
    return False


def copy_with_progress(src, show_progress, interactive):
    """用 backup API 复制；show_progress 时显示百分比并允许按 n 取消。"""
    tmp = os.path.join(tempfile.gettempdir(),
                       "dbview-%s-%s-%d.db" % (
                           os.path.splitext(os.path.basename(src))[0],
                           datetime.now().strftime("%H%M%S"), os.getpid()))
    src_conn = None
    last_err = None
    for ro_uri in (uri_of(src, True, False), uri_of(src, True, True)):
        try:
            src_conn = sqlite3.connect(ro_uri, uri=True, timeout=20.0)
            src_conn.execute("PRAGMA query_only=1")
            break
        except Exception as e:                       # noqa: BLE001
            last_err = e
            src_conn = None
    if src_conn is None:
        raise RuntimeError("打不开源库：%s" % last_err)

    try:
        page_size = int(src_conn.execute("PRAGMA page_size").fetchone()[0]) or 4096
    except Exception:
        page_size = 4096
    dst_conn = sqlite3.connect(tmp)
    cancelled = {"v": False}
    t0 = time.time()
    last_paint = [0.0]

    def on_progress(status, remaining, total):
        if interactive and _kbhit_cancel():
            cancelled["v"] = True
            raise CopyCancelled()
        if not show_progress:
            return
        now = time.time()
        if now - last_paint[0] < 0.3 and remaining:
            return
        last_paint[0] = now
        done = (total - remaining) if total else 0
        pct = (done * 100.0 / total) if total else 100.0
        el = now - t0
        rate = done / el if el > 0.2 and done else 0
        eta = (remaining / rate) if rate else 0
        print("\r复制中 %5.1f%%  %s/%s  已用 %s  剩余 ~%s  （输 n 取消）"
              % (pct, _fmt_mb(done * page_size),
                 _fmt_mb((total or 0) * page_size),
                 _hms(el), _hms(eta)) + "   ", end="", flush=True)

    try:
        src_conn.backup(dst_conn, pages=512, progress=on_progress, sleep=0.02)
    except CopyCancelled:
        dst_conn.close()
        src_conn.close()
        try:
            os.remove(tmp)
        except OSError:
            pass
        raise
    finally:
        try:
            dst_conn.close()
        except Exception:
            pass
        try:
            src_conn.close()
        except Exception:
            pass
    if show_progress:
        print("\r" + " " * 88 + "\r", end="", flush=True)
    return tmp


def space_check(path):
    """返回 (是否够, 需要字节, 可用字节)。"""
    size = 0
    try:
        size = os.path.getsize(path)
        for suffix in ("-wal", "-shm"):
            extra = path + suffix
            if os.path.exists(extra):
                size += os.path.getsize(extra)
    except OSError:
        return True, 0, 0
    need = int(size * 1.25) + 64 * 1048576
    free = shutil.disk_usage(tempfile.gettempdir()).free
    return free >= need, need, free


def ask_space_choice(need, free):
    print("[警告] 临时目录空间不足：需要约 %s，可用 %s"
          % (_fmt_mb(need), _fmt_mb(free)))
    print("   1) 跳过优化，直接用这个库导出（不复制，较慢）")
    print("   2) 取消，换个库或清理临时目录后再试")
    try:
        ans = input("请输入序号（默认 1）：").strip()
    except (EOFError, KeyboardInterrupt):
        return 1
    return 2 if ans == "2" else 1


def optimize_db(conn):
    """在给定库上建索引 + ANALYZE，返回创建的对象清单（用于回滚）。"""
    created = []
    # 只对"本项目的表结构"加索引，别的库不乱动
    try:
        cols = [r[1] for r in conn.execute("PRAGMA table_info(evidence)")]
    except Exception:
        cols = []
    if {"persona_id", "source", "state"} <= set(cols):
        try:
            conn.execute("CREATE INDEX IF NOT EXISTS ix_exp_lookup"
                         " ON evidence(persona_id, source, state)")
            created.append(("index", "ix_exp_lookup"))
        except Exception:
            pass
    had_stat = False
    try:
        had_stat = bool(conn.execute(
            "SELECT count(*) FROM sqlite_master WHERE name='sqlite_stat1'"
        ).fetchone()[0])
    except Exception:
        pass
    try:
        conn.execute("ANALYZE")
        if not had_stat:
            created.append(("table", "sqlite_stat1"))
    except Exception:
        pass
    conn.commit()
    return created


def undo_optimize(conn, created):
    for kind, name in created:
        try:
            if kind == "index":
                conn.execute('DROP INDEX IF EXISTS "%s"' % name)
            else:
                conn.execute('DROP TABLE IF EXISTS "%s"' % name)
        except Exception:
            pass
    try:
        conn.commit()
    except Exception:
        pass


def count_rows(db_path, obj, timeout=5.0):
    """独立连接里数行数（SQLite 连接不能跨线程）。拿不到返回 None。"""
    box = {}

    def work():
        try:
            c2, _ = open_ro(db_path)
            if c2:
                try:
                    box["n"] = c2.execute(
                        'SELECT COUNT(*) FROM "%s"'
                        % obj.replace('"', '""')).fetchone()[0]
                finally:
                    c2.close()
        except Exception:
            pass

    t = threading.Thread(target=work, daemon=True)
    t.start()
    t.join(timeout)
    return box.get("n")


# ---------------------------------------------------------------- 选择对象
def list_objects(conn, mode):
    views = [r["name"] for r in conn.execute(
        "SELECT name FROM sqlite_master WHERE type='view' ORDER BY name")]
    tables = [r["name"] for r in conn.execute(
        "SELECT name FROM sqlite_master WHERE type='table'"
        " AND name NOT LIKE 'sqlite_%' ORDER BY name")]
    if mode == "views":
        return views, "view"
    if mode == "tables":
        return tables, "table"
    return views + tables, "object"


def order_objects(db_path, objs, prefer_preset=True):
    """本项目的 8 个视图按预设顺序；否则按行数升序（数不出排最后）。"""
    sizes = {}
    for o in objs:
        sizes[o] = count_rows(db_path, o, 5.0)
    if prefer_preset and all(o in PRESET_BANLIST for o in objs):
        ordered = [o for o in PRESET_BANLIST if o in objs]
        return ordered, sizes
    known = [o for o in objs if sizes[o] is not None]
    known.sort(key=lambda o: sizes[o])
    unknown = [o for o in objs if sizes[o] is None]
    return known + unknown, sizes


def choose(objs, kind):
    print("")
    print("发现 %d 个%s：" % (len(objs), "视图" if kind == "view" else kind))
    for i, o in enumerate(objs, 1):
        print("  %2d) %s" % (i, o))
    print("   0) 全部导出（按行数从小到大）")
    print("")
    while True:
        try:
            raw = input("请输入序号（直接回车 = 全部）：").strip()
        except (EOFError, KeyboardInterrupt):
            return None, False
        if raw in ("", "0"):
            return list(objs), True
        if raw.isdigit() and 1 <= int(raw) <= len(objs):
            return [objs[int(raw) - 1]], False
        print("  序号无效，请重新输入。")


# ---------------------------------------------------------------- 输出目录
def make_out_dir(name, stamp):
    bases = []
    for b in (desktop_dir(),
              os.path.join(os.path.expanduser("~"), "Documents"),
              os.path.join(os.path.expanduser("~"), "Documents", "视图导出"),
              os.getcwd(), tempfile.gettempdir()):
        if b and os.path.isdir(b) and b not in bases:
            bases.append(b)
    for base in bases:
        target = os.path.join(base, "%s-%s" % (name, stamp))
        for i in range(1, 6):
            p = target if i == 1 else "%s-%d" % (target, i)
            try:
                os.makedirs(p)
                return p, base
            except FileExistsError:
                continue
            except OSError:
                break
    return None, None


# ---------------------------------------------------------------- 导出
def progress_text(st):
    el = time.time() - st["t0"]
    nv = st["nviews"]
    if st.get("error"):
        return "[%d/%d] %-14s 失败：%s" % (st["idx"], nv, st["obj"],
                                          str(st["error"])[:40])
    if st["rows"] == 0:
        note = "仍在计算（对象较慢）" if el > 10 else "正在导出…"
        return "[%d/%d] %-14s %s  已用 %s" % (st["idx"], nv, st["obj"], note,
                                              _hms(el))
    total = st.get("total")
    n = st["rows"]
    if total:
        pct = min(100.0, n * 100.0 / total)
        rate = n / el if el > 0.2 else 0
        eta = (total - n) / rate if rate > 0 else 0
        s = "[%d/%d] %-14s %s/%s %5.1f%%  已用 %s  剩余 ~%s" % (
            st["idx"], nv, st["obj"], format(n, ","), format(total, ","),
            pct, _hms(el), _hms(eta))
    else:
        s = "[%d/%d] %-14s %s 行  已用 %s" % (
            st["idx"], nv, st["obj"], format(n, ","), _hms(el))
    return s.ljust(78)


def calc_widths(conn, obj, cols, sample=200, cap=60):
    widths = [max(8, disp_width(c)) for c in cols]
    try:
        cur = conn.execute('SELECT * FROM "%s" LIMIT %d'
                           % (obj.replace('"', '""'), sample))
        for r in cur.fetchall():
            for i, c in enumerate(cols):
                v = r[c]
                if v is None:
                    continue
                w = disp_width(str(v))
                if w > widths[i]:
                    widths[i] = min(w, cap)
    except Exception:
        pass
    return [max(10, min(w + 2, cap)) for w in widths]


def export_one(conn, obj, out_dir, ext, sizes, idx, nviews):
    safe = "".join("_" if c in '\\/:*?"<>|' else c for c in obj)
    path = os.path.join(out_dir, "%s%s" % (safe, ext))
    st = {"idx": idx, "obj": obj, "rows": 0, "total": sizes.get(obj),
          "t0": time.time(), "done": False, "error": None, "nviews": nviews}
    stop = threading.Event()

    def beat():
        while not stop.wait(0.5):
            print("\r" + progress_text(st) + "        ", end="", flush=True)

    print("")
    th = threading.Thread(target=beat, daemon=True)
    th.start()
    try:
        cur = conn.execute('SELECT * FROM "%s"' % obj.replace('"', '""'))
        cols = [d[0] for d in cur.description]
        n = 0
        if ext == ".xlsx":
            widths = calc_widths(conn, obj, cols)
            wb = Workbook(write_only=True)
            ws = wb.create_sheet("数据")
            for i, c in enumerate(cols, 1):
                ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
            ws.freeze_panes = "A2"
            head = []
            for c in cols:
                cell = WriteOnlyCell(ws, value=c)
                cell.font = Font(bold=True)
                head.append(cell)
            ws.append(head)
            while True:
                rows = cur.fetchmany(2000)
                if not rows:
                    break
                for r in rows:
                    ws.append([cell_value(r[c]) for c in cols])
                n += len(rows)
                st["rows"] = n
            wb.save(path)
        else:
            with open(path, "w", encoding="utf-8-sig", newline="") as f:
                w = csv.writer(f)
                w.writerow(cols)
                while True:
                    rows = cur.fetchmany(2000)
                    if not rows:
                        break
                    for r in rows:
                        w.writerow(["" if cell_value(r[c]) is None
                                    else cell_value(r[c]) for c in cols])
                    n += len(rows)
                    st["rows"] = n
        st["done"] = True
        return True, n, len(cols), None
    except Exception as e:                           # noqa: BLE001
        return False, 0, 0, str(e)
    finally:
        stop.set()
        th.join(0.3)
        print("\r" + progress_text(st) + ("  完成" if st["done"] else "")
              + "        ", flush=True)


# ---------------------------------------------------------------- 主流程
def run_export(db_path, args):
    db_path = os.path.abspath(db_path)
    name = os.path.splitext(os.path.basename(db_path))[0]
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    out_dir, base = make_out_dir(name, stamp)
    if not out_dir:
        print("[错误] 桌面、文档、当前目录、临时目录都写不进去。")
        return 1
    if base != desktop_dir():
        print("[提示] 桌面写不进去，本次输出改到：%s" % base)

    print("")
    print("数据库：%s" % db_path)
    print("输出到：%s" % out_dir)

    # 1) 先只读看一眼结构
    conn, err = open_ro(db_path)
    if conn is None:
        print("[错误] 打不开数据库：%s" % err)
        print("       若提示 malformed，说明是运行中直接复制的坏副本；")
        print("       请在主程序里用备份功能重新做一份。")
        return 1
    if args.check:
        try:
            info = conn.execute("PRAGMA quick_check").fetchone()[0]
            print("完整性检查：%s" % info)
        except Exception as e:                       # noqa: BLE001
            print("完整性检查失败：%s" % e)

    objs, kind = list_objects(conn, args.objects)
    if not objs:
        print("这个数据库里没有可导出的%s。" % ("视图" if kind == "view" else "对象"))
        conn.close()
        return 0

    if args.yes:
        chosen, all_mode = list(objs), True
    else:
        chosen, all_mode = choose(objs, kind)
        if not chosen:
            conn.close()
            return 1

    # 2) 判断要不要优化
    work_db = db_path
    temp_copy = None
    created = []
    conn.close()

    size_mb = os.path.getsize(db_path) / 1048576.0
    interactive = sys.stdin.isatty()
    want_opt = bool(args.optimize)

    if not args.optimize_inplace and not want_opt and interactive \
            and size_mb > 30:
        try:
            ans = input("\n这个库有 %.0f MB，要不要在临时副本上建索引/统计以加速？"
                        "（不改原文件）(Y/n) " % size_mb).strip().lower()
        except (EOFError, KeyboardInterrupt):
            ans = "n"
        want_opt = ans in ("", "y", "yes")

    if args.optimize_inplace:
        print("")
        print("[警告] --optimize-inplace 会直接修改这个 .db：")
        print("        会建索引、跑 ANALYZE（导出结束会删除自己建的索引）。")
        print("        请确保这是【备份副本】，不是正在运行的库。")
        try:
            ok = input("确认继续请输入大写 YES：").strip()
        except (EOFError, KeyboardInterrupt):
            ok = ""
        if ok != "YES":
            print("已取消。")
            return 1
        side_before = snapshot_side(db_path)
        wconn = open_rw(db_path)
        created = optimize_db(wconn)
        wconn.close()
        print("已优化原库：%s" % (", ".join("%s %s" % c for c in created) or "无"))
        work_db = db_path
    elif want_opt:
        enough, need, free = space_check(db_path)
        if not enough:
            choice = (ask_space_choice(need, free) if interactive else 1)
            if choice != 1:
                print("已取消，未做任何改动。")
                return 1
            print("已选择：跳过优化，直接用原库导出（只读）。")
            want_opt = False
        if want_opt:
            show_progress = (size_mb >= 1024)      # 1 GB 以上才显示复制进度
            print("\n正在复制临时副本（%s）%s"
                  % (_fmt_mb(size_mb * 1048576),
                     "…按提示可输 n 取消" if show_progress else "…"))
            try:
                temp_copy = copy_with_progress(db_path, show_progress,
                                               interactive)
            except CopyCancelled:
                print("已取消复制，改用原库导出（只读，不会删除你的库）。")
                temp_copy = None
            except Exception as e:                   # noqa: BLE001
                print("[警告] 复制失败，改用原库导出：%s" % e)
                temp_copy = None
            if temp_copy:
                try:
                    wconn = open_rw(temp_copy)
                    created = optimize_db(wconn)
                    wconn.close()
                    print("已在临时副本上优化：%s"
                          % (", ".join("%s %s" % c for c in created) or "无"))
                    work_db = temp_copy
                except Exception as e:               # noqa: BLE001
                    print("[警告] 优化临时副本失败，仍用它导出：%s" % e)
                    work_db = temp_copy

    # 3) 用（可能优化过的）库导出
    conn, err = open_ro(work_db)
    if conn is None:
        print("[错误] 优化副本打不开：%s" % err)
        return 1

    if not all_mode:
        order = chosen
        sizes = ({o: None for o in chosen} if args.no_count
                 else {o: count_rows(work_db, o, 5.0) for o in chosen})
    elif args.no_count:
        order = chosen
        sizes = {o: None for o in chosen}
    else:
        order, sizes = order_objects(work_db, chosen, prefer_preset=True)

    ext = ".csv" if args.format == "csv" else (".xlsx" if HAS_XLSX else ".csv")
    manifest = ["数据库：%s" % db_path,
                "导出时间：%s" % datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                "对象类型：%s" % ("视图" if kind == "view" else kind),
                "对象数量：%d" % len(order),
                "格式：%s" % ext,
                ""]
    if temp_copy:
        manifest.append("本次在临时副本上优化后导出，原文件未修改。")
        manifest.append("")

    ok_count = bad_count = 0
    for idx, obj in enumerate(order, 1):
        ok, n, ncol, err = export_one(conn, obj, out_dir, ext, sizes,
                                      idx, len(order))
        if ok:
            manifest.append("[OK ] %s\t%d 行\t%d 列" % (obj, n, ncol))
            ok_count += 1
        else:
            manifest.append("[BAD] %s\t%s" % (obj, err))
            bad_count += 1
    conn.close()

    # 4) 收尾：删临时副本、回滚原库改动、清 side 文件
    if temp_copy:
        try:
            os.remove(temp_copy)
            print("临时副本已删除。")
        except OSError as e:
            print("临时副本没删掉（%s）：%s" % (e, temp_copy))
    if args.optimize_inplace and created:
        try:
            wconn = open_rw(db_path)
            undo_optimize(wconn, created)
            wconn.close()
            print("已回滚原库上创建的索引/统计。")
        except Exception as e:                       # noqa: BLE001
            print("[警告] 回滚失败：%s" % e)
    if args.optimize_inplace:
        cleanup_side(db_path, side_before)

    manifest += ["", "成功 %d 个，失败 %d 个" % (ok_count, bad_count)]
    with open(os.path.join(out_dir, "清单.txt"), "w",
              encoding="utf-8-sig") as f:
        f.write("\n".join(manifest) + "\n")

    print("")
    print("完成：成功 %d，失败 %d" % (ok_count, bad_count))
    print("文件都在：%s" % out_dir)
    return 0


def ask_path():
    print("=" * 66)
    print("DB 视图导出工具")
    print("=" * 66)
    print("请把 .db（或它的快捷方式、或包含 .db 的文件夹）拖到这个窗口，")
    print("也可以直接粘贴路径，然后回车：")
    while True:
        try:
            raw = input("> ")
        except (EOFError, KeyboardInterrupt):
            return None
        path = clean_path(raw)
        if not path:
            continue
        db, err = resolve_input(path)
        if db:
            return db
        print("  %s：%s" % (err or "无法识别", path))


def parse_args(argv):
    ap = argparse.ArgumentParser(add_help=True)
    ap.add_argument("--db", help="数据库路径")
    ap.add_argument("--objects", choices=["views", "tables", "all"],
                    default="views", help="导出视图/表/两者（默认视图）")
    ap.add_argument("--format", choices=["xlsx", "csv"], default="xlsx")
    ap.add_argument("--yes", action="store_true", help="不交互，全部导出")
    ap.add_argument("--no-count", action="store_true", help="不预统计行数")
    ap.add_argument("--check", action="store_true", help="先跑 quick_check")
    ap.add_argument("--optimize", action="store_true",
                    help="在临时副本上建索引/ANALYZE 后导出（不改原库）")
    ap.add_argument("--optimize-inplace", action="store_true",
                    help="直接优化原库（会改文件，需输 YES，请用备份库）")
    return ap.parse_args(argv)


def _frozen():
    return bool(getattr(sys, "frozen", False))


def _write_log(name, text):
    try:
        path = os.path.join(tempfile.gettempdir(), name)
        with open(path, "a", encoding="utf-8") as f:
            f.write("%s %s\n" % (datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                                 text))
        return path
    except Exception:
        return ""


def main(argv=None):
    args = parse_args(argv if argv is not None else sys.argv[1:])
    while True:
        db = args.db
        if not db:
            db = ask_path()
        else:
            resolved, err = resolve_input(db)
            if not resolved:
                print("无法使用：%s（%s）" % (db, err))
                db = ask_path()
            else:
                db = resolved
        if not db:
            return 1
        code = run_export(db, args)
        if args.yes or args.db:
            return code
        print("")
        try:
            again = input("还要导出别的库吗？(y/N) ").strip().lower()
        except (EOFError, KeyboardInterrupt):
            return code
        if again != "y":
            return code
        args.db = None


if __name__ == "__main__":
    import traceback
    _write_log("DB视图导出-启动日志.txt",
               "启动 argv=%s frozen=%s" % (sys.argv, _frozen()))
    code = 0
    try:
        code = main()
    except SystemExit as e:
        code = e.code if isinstance(e.code, int) else 0
    except BaseException:                            # noqa: BLE001
        tb = traceback.format_exc()
        print(tb)
        p1 = _write_log("DB视图导出-错误.txt", tb)
        p2 = ""
        try:
            p2 = os.path.join(desktop_dir(), "DB视图导出-错误.txt")
            with open(p2, "w", encoding="utf-8") as f:
                f.write(tb)
        except Exception:
            p2 = ""
        print("")
        print("出错了，错误详情已保存到：")
        if p1:
            print("  " + p1)
        if p2:
            print("  " + p2)
        code = 1
    finally:
        if _frozen():
            try:
                input("\n按回车退出…")
            except Exception:
                pass
    sys.exit(code)
